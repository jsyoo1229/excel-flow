import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font


def clean_csv_files(data_folder, output_file, remove_missing=True):
    # Find all CSV files in the selected folder
    csv_files = list(data_folder.glob("*.csv"))

    if not csv_files:
        raise ValueError("No CSV files were found in the selected folder.")


    # Read CSV files
    dataframes = []
    failed_files = []

    for file in csv_files:
        try:
            df = pd.read_csv(file)

            # Remove whitespace from column names
            df.columns = df.columns.str.strip()

            dataframes.append((file, df))

        except Exception:
            failed_files.append(file)


    if not dataframes:
        raise ValueError("None of the CSV files could be read.")


    # Count column structures
    column_counts = {}

    for file, df in dataframes:
        columns = tuple(df.columns)

        if columns in column_counts:
            column_counts[columns] += 1
        else:
            column_counts[columns] = 1


    # Use the most common column structure
    main_columns = max(
        column_counts,
        key=column_counts.get
    )


    # Keep files matching the main structure
    valid_dataframes = []
    skipped_files = []

    for file, df in dataframes:
        if tuple(df.columns) == main_columns:
            valid_dataframes.append(df)
        else:
            skipped_files.append(file)


    # Combine valid CSV files
    combined = pd.concat(
        valid_dataframes,
        ignore_index=True
    )


    # Clean text columns
    for column in combined.select_dtypes(include="str").columns:
        combined[column] = combined[column].str.strip()
        combined[column] = combined[column].replace("", pd.NA)


    # Find text-like columns
    string_columns = (
        combined
        .select_dtypes(include="str")
        .columns
        .tolist()
    )

    object_columns = (
        combined
        .select_dtypes(
            include="object",
            exclude="str"
        )
        .columns
        .tolist()
    )

    candidate_columns = string_columns + object_columns


    # Convert numeric-looking columns
    for column in candidate_columns:
        cleaned_values = (
            combined[column]
            .astype("string")
            .str.replace(",", "", regex=False)
            .str.replace("$", "", regex=False)
        )

        numeric_values = pd.to_numeric(
            cleaned_values,
            errors="coerce"
        )

        non_missing = combined[column].notna()

        if numeric_values[non_missing].notna().all():
            combined[column] = numeric_values


    # Detect date columns
    for column in combined.select_dtypes(include="str").columns:
        try:
            parsed_dates = pd.to_datetime(
                combined[column],
                format="mixed",
                errors="coerce"
            )

            non_missing = combined[column].notna()

            if parsed_dates[non_missing].notna().all():
                combined[column] = parsed_dates

        except Exception:
            pass


    # Remove duplicates
    rows_before_cleaning = len(combined)

    cleaned = combined.drop_duplicates()

    duplicate_rows_removed = (
        rows_before_cleaning - len(cleaned)
    )


    # Handle missing values
    rows_before_missing = len(cleaned)

    if remove_missing:
        cleaned = cleaned.dropna()

    missing_rows_removed = (
        rows_before_missing - len(cleaned)
    )


    # Export to Excel
    cleaned.to_excel(
        output_file,
        index=False
    )


    # Format the Excel output
    workbook = load_workbook(output_file)
    worksheet = workbook.active


    # Make header cells bold
    for cell in worksheet[1]:
        cell.font = Font(bold=True)


    # Freeze the header row
    worksheet.freeze_panes = "A2"


    # Add filters to the header row
    worksheet.auto_filter.ref = worksheet.dimensions


        # Format date columns
    for column_index, column_name in enumerate(
        cleaned.columns,
        start=1
    ):
        if pd.api.types.is_datetime64_any_dtype(
            cleaned[column_name]
        ):
            for row_index in range(
                2,
                worksheet.max_row + 1
            ):
                worksheet.cell(
                    row=row_index,
                    column=column_index
                ).number_format = "yyyy-mm-dd"


    # Automatically adjust column widths
    for column_cells in worksheet.columns:
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                cell_length = len(str(cell.value))

                max_length = max(
                    max_length,
                    cell_length
                )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[column_letter].width = (
            min(max_length + 2, 50)
        )


    # Save the formatted Excel file
    workbook.save(output_file)


    # Return cleaning results
    return {
        "files_found": len(csv_files),
        "files_read": len(dataframes),
        "files_processed": len(valid_dataframes),
        "files_skipped": len(skipped_files),
        "files_failed": len(failed_files),
        "rows_before_cleaning": rows_before_cleaning,
        "duplicate_rows_removed": duplicate_rows_removed,
        "missing_rows_removed": missing_rows_removed,
        "rows_after_cleaning": len(cleaned),
    }